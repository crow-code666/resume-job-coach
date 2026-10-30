
import io
from pathlib import Path

import pdfplumber
from docx import Document


class ResumeParser:
    """Extract readable text from PDF and DOCX resumes."""

    SUPPORTED_FORMATS = {".pdf", ".docx"}

    def extract_text(self, uploaded_file) -> str:
        """
        Extract text from an uploaded PDF or DOCX file.

        Args:
            uploaded_file: Streamlit UploadedFile or file-like object.

        Returns:
            Extracted text as a string.

        Raises:
            ValueError: If the file format is unsupported or text is empty.
        """
        filename = getattr(uploaded_file, "name", "resume")
        extension = Path(filename).suffix.lower()

        if extension not in self.SUPPORTED_FORMATS:
            raise ValueError(
                "Unsupported file format. Please upload a PDF or DOCX file."
            )

        file_bytes = uploaded_file.getvalue()

        if not file_bytes:
            raise ValueError("The uploaded file is empty.")

        if extension == ".pdf":
            text = self._extract_pdf(file_bytes)
        else:
            text = self._extract_docx(file_bytes)

        text = self._clean_text(text)

        if not text.strip():
            raise ValueError(
                "No readable text was found. "
                "The PDF may be scanned or image-based."
            )

        return text

    @staticmethod
    def _extract_pdf(file_bytes: bytes) -> str:
        """Extract text from all pages of a PDF."""
        extracted_pages = []

        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    extracted_pages.append(page_text)

        return "\n\n".join(extracted_pages)

    @staticmethod
    def _extract_docx(file_bytes: bytes) -> str:
        """Extract paragraphs and table text from a DOCX file."""
        document = Document(io.BytesIO(file_bytes))
        extracted_parts = []

        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                extracted_parts.append(paragraph.text)

        for table in document.tables:
            for row in table.rows:
                row_text = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                ]

                if row_text:
                    extracted_parts.append(" | ".join(row_text))

        return "\n".join(extracted_parts)

    @staticmethod
    def _clean_text(text: str) -> str:
        """Normalize whitespace without removing useful line breaks."""
        lines = []

        for line in text.splitlines():
            cleaned_line = " ".join(line.split())
            if cleaned_line:
                lines.append(cleaned_line)

        return "\n".join(lines)

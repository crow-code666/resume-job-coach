
import io
import zipfile
from pathlib import Path

import pdfplumber
from docx import Document


class ResumeParser:
    """Extract readable text from PDF and DOCX resumes."""

    SUPPORTED_FORMATS = {".pdf", ".docx"}
    MAX_FILE_SIZE_BYTES = 20 * 1024 * 1024
    MAX_EXTRACTED_TEXT_CHARS = 100_000
    MAX_DOCX_EXPANDED_SIZE_BYTES = 100 * 1024 * 1024
    MAX_DOCX_PARTS = 10_000

    def extract_text(self, uploaded_file) -> str:
        """
        Extract text from an uploaded PDF or DOCX file.

        Args:
            uploaded_file: Streamlit UploadedFile or file-like object.

        Returns:
            Extracted text as a string.

        Raises:
            ValueError: If the file format is unsupported or text is empty.
        """
        filename = getattr(uploaded_file, "name", "resume")
        extension = Path(filename).suffix.lower()

        if extension not in self.SUPPORTED_FORMATS:
            raise ValueError(
                "Unsupported file format. Please upload a PDF or DOCX file."
            )

        file_bytes = uploaded_file.getvalue()

        if not file_bytes:
            raise ValueError("The uploaded file is empty.")
        if len(file_bytes) > self.MAX_FILE_SIZE_BYTES:
            raise ValueError("The uploaded file is too large. The limit is 20 MB.")

        if extension == ".pdf":
            text = self._extract_pdf(file_bytes)
        else:
            self._validate_docx_container(file_bytes)
            text = self._extract_docx(file_bytes)

        text = self._clean_text(text)

        if not text.strip():
            raise ValueError(
                "No readable text was found. "
                "The PDF may be scanned or image-based."
            )
        if len(text) > self.MAX_EXTRACTED_TEXT_CHARS:
            raise ValueError(
                "The extracted resume text exceeds the 100,000 character limit."
            )

        return text

    @classmethod
    def _validate_docx_container(cls, file_bytes: bytes) -> None:
        """Reject malformed or implausibly large DOCX archives before parsing."""
        try:
            with zipfile.ZipFile(io.BytesIO(file_bytes)) as archive:
                parts = archive.infolist()
                expanded_size = sum(part.file_size for part in parts)
        except zipfile.BadZipFile as error:
            raise ValueError("The DOCX file is invalid or damaged.") from error

        if len(parts) > cls.MAX_DOCX_PARTS:
            raise ValueError("The DOCX file contains too many internal parts.")
        if expanded_size > cls.MAX_DOCX_EXPANDED_SIZE_BYTES:
            raise ValueError("The DOCX file expands beyond the 100 MB safety limit.")

    @staticmethod
    def _extract_pdf(file_bytes: bytes) -> str:
        """Extract text from all pages of a PDF."""
        extracted_pages = []

        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    extracted_pages.append(page_text)

        return "\n\n".join(extracted_pages)

    @staticmethod
    def _extract_docx(file_bytes: bytes) -> str:
        """Extract paragraphs and table text from a DOCX file."""
        document = Document(io.BytesIO(file_bytes))
        extracted_parts = []

        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                extracted_parts.append(paragraph.text)

        for table in document.tables:
            for row in table.rows:
                row_text = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                ]

                if row_text:
                    extracted_parts.append(" | ".join(row_text))

        return "\n".join(extracted_parts)

    @staticmethod
    def _clean_text(text: str) -> str:
        """Normalize whitespace without removing useful line breaks."""
        lines = []

        for line in text.splitlines():
            cleaned_line = " ".join(line.split())
            if cleaned_line:
                lines.append(cleaned_line)

        return "\n".join(lines)
